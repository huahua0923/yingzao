# -*- coding: utf-8 -*-
"""docx → 结构化文本（**判据用**，不是给人读的）：段落文本序列 + 表格单元格文本序列。

用法: python tools/docx2txt.py <in.docx> <out.json>

★ 用途：两份 docx 的**结构对账** —— 生成器换了、或 md 改过之后，把新旧两份各抽一遍，
  逐条比；除有意改动外必须逐字相同。**这是「生成器忠实」这条判据的量具**。
★ 盲区（必须声明）：只比**文本**，不比字体/字号/加粗/表格样式 ⇒ **版面差异它一概看不见**。
  版面另由 `md2docx.convert()` 里写死的版面参数保证（A4 / 边距 / Table Grid）。
"""
import json
import sys

import docx


def dump(path):
    d = docx.Document(path)
    paras = [p.text for p in d.paragraphs]
    tables = [[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
    return {"file": path, "n_para": len(paras), "n_table": len(tables),
            "paragraphs": paras, "tables": tables}


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("用法: python tools/docx2txt.py <in.docx> <out.json>"); sys.exit(2)
    obj = dump(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8", newline="\n") as f:   # ★ 字节层口径：LF
        json.dump(obj, f, ensure_ascii=False, indent=1)
    print("%s → %s  段落 %d 表格 %d" % (sys.argv[1], sys.argv[2], obj["n_para"], obj["n_table"]))
